"""
Export Excel DETERMINISTE des resultats de provisionnement.

Genere un classeur mis en forme (sans aucun recours a un LLM : chaque cellule
est calculee et tracable) a partir du resultat de consolider_best_estimate().

Feuilles :
    - Synthese        : Best Estimate par segment + total, vs SAP comptable reel
    - <SEGMENT>        : comparatif detaille des 4 methodes (une feuille/segment)
    - Parametres       : date d'arrete, hypotheses, provenance
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Chartre graphique du classeur (aligne sur la palette plateforme) ---
BLEU = "2A78D6"
BLEU_FONCE = "184F95"
GRIS_ENTETE = "F0EFEC"
VERT = "0CA30C"      # boni
ROUGE = "D03B3B"     # mali
INK = "0B0B0B"

_FMT_DT = '#,##0 "DT"'
_FMT_PCT = '0.0"%"'

_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_titre(ws, cellule: str, texte: str, taille: int = 14):
    ws[cellule] = texte
    ws[cellule].font = Font(bold=True, size=taille, color=BLEU_FONCE)


def _ecrire_entete(ws, ligne: int, colonnes: list[str], col_debut: int = 1):
    for j, nom in enumerate(colonnes):
        c = ws.cell(row=ligne, column=col_debut + j, value=nom)
        c.font = Font(bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=GRIS_ENTETE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER


def _ecrire_ligne(ws, ligne: int, valeurs: list, formats: list[str | None], col_debut: int = 1):
    for j, (val, fmt) in enumerate(zip(valeurs, formats)):
        c = ws.cell(row=ligne, column=col_debut + j, value=val)
        c.border = _BORDER
        if fmt:
            c.number_format = fmt
        if j == 0:
            c.font = Font(bold=False)


def _auto_width(ws, largeurs: dict[int, int]):
    for col, w in largeurs.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _feuille_synthese(wb: Workbook, conso: dict, date_arrete: date):
    ws = wb.active
    ws.title = "Synthese"
    ws.sheet_view.showGridLines = False

    _style_titre(ws, "A1", "AutoReserve — Synthese Best Estimate", 16)
    ws["A2"] = f"Provisions techniques sinistres auto — arrete au {date_arrete:%d/%m/%Y}"
    ws["A2"].font = Font(italic=True, color="52514E")

    syn = conso["synthese"]
    total = conso["total"]

    cols = ["Segment", "Methode retenue", "Paye cumule", "Charge ultime",
            "Reserves estimees", "SAP comptable reel", "IBNR estime", "Boni / Mali"]
    fmts = [None, None, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT]

    ligne = 4
    _ecrire_entete(ws, ligne, cols)
    ligne += 1
    for seg, row in syn.iterrows():
        vals = [seg, row["methode_retenue"], row["paye"], row["ultime"],
                row["reserves_estimees"], row["sap_comptable"],
                row["ibnr_estime"], row["boni_mali"]]
        _ecrire_ligne(ws, ligne, vals, fmts)
        # couleur boni/mali
        c = ws.cell(row=ligne, column=8)
        c.font = Font(bold=True, color=VERT if row["boni_mali"] >= 0 else ROUGE)
        ligne += 1

    # ligne TOTAL
    vals = ["TOTAL", "", syn["paye"].sum(), syn["ultime"].sum(),
            total["reserves_estimees"], total["sap_comptable"],
            total["ibnr_estime"], total["boni_mali"]]
    _ecrire_ligne(ws, ligne, vals, fmts)
    for j in range(1, 9):
        cc = ws.cell(row=ligne, column=j)
        cc.font = Font(bold=True, color=INK)
        cc.fill = PatternFill("solid", fgColor=GRIS_ENTETE)
    ws.cell(row=ligne, column=8).font = Font(
        bold=True, color=VERT if total["boni_mali"] >= 0 else ROUGE)

    ligne += 2
    ws.cell(row=ligne, column=1,
            value="Boni (+) = sur-provision comptable · Mali (-) = sous-provision. "
                  "Montants en dinars tunisiens (DT).").font = Font(italic=True, color="898781")

    _auto_width(ws, {1: 14, 2: 22, 3: 16, 4: 16, 5: 18, 6: 18, 7: 16, 8: 16})


def _feuille_segment(wb: Workbook, nature: str, tab: pd.DataFrame):
    ws = wb.create_sheet(title=nature.capitalize())
    ws.sheet_view.showGridLines = False
    _style_titre(ws, "A1", f"Comparatif des methodes — {nature}", 14)

    cols = ["Methode", "Paye", "Ultime", "Reserves", "SAP comptable",
            "IBNR estime", "Boni / Mali"]
    fmts = [None, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT, _FMT_DT]

    ligne = 3
    _ecrire_entete(ws, ligne, cols)
    ligne += 1
    for methode, row in tab.iterrows():
        vals = [methode, row["paye"], row["ultime"], row["reserves"],
                row["sap_comptable"], row["ibnr_estime"], row["boni_mali"]]
        _ecrire_ligne(ws, ligne, vals, fmts)
        c = ws.cell(row=ligne, column=7)
        c.font = Font(bold=True, color=VERT if row["boni_mali"] >= 0 else ROUGE)
        ligne += 1

    _auto_width(ws, {1: 24, 2: 16, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16})


def _feuille_parametres(wb: Workbook, conso: dict, date_arrete: date):
    ws = wb.create_sheet(title="Parametres")
    ws.sheet_view.showGridLines = False
    _style_titre(ws, "A1", "Parametres & hypotheses", 14)

    lignes = [
        ("Date d'arrete", f"{date_arrete:%d/%m/%Y}"),
        ("Segmentation", "CORPOREL / MATERIEL"),
        ("Methode retenue (defaut)", "Chain-Ladder ecrete"),
        ("Ecretement des graves", "Capping des paiements au quantile 99%"),
        ("Benchmark de controle", "SAP comptable (diagonale courante)"),
        ("Devise", "Dinar tunisien (DT)"),
        ("Genere par", "AutoReserve (calculs deterministes, sans IA generative)"),
    ]
    for i, (k, v) in enumerate(lignes, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color=INK)
        ws.cell(row=i, column=2, value=v)
    _auto_width(ws, {1: 28, 2: 55})


def exporter_best_estimate(conso: dict, date_arrete: date | None = None) -> bytes:
    """
    Construit le classeur Excel de synthese et renvoie ses octets
    (prets pour st.download_button ou ecriture sur disque).
    """
    if date_arrete is None:
        date_arrete = date.today()

    wb = Workbook()
    _feuille_synthese(wb, conso, date_arrete)
    for seg, tab in conso["comparatifs_par_segment"].items():
        _feuille_segment(wb, seg, tab)
    _feuille_parametres(wb, conso, date_arrete)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
